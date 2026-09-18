from copy import deepcopy
from datetime import datetime
from pathlib import Path
import json
import tempfile
import unittest
from unittest.mock import MagicMock, patch

from profiles import (adjust_range, daily_basal, differences, excel_sheets, export_excel, import_excel, load_profiles,
                      minute, normalize_schedule, overwrite_profile, profile_changes, save_profile, validate_profile, value_at)

ROOT = Path(__file__).resolve().parents[1]


class ProfileTests(unittest.TestCase):
    def setUp(self):
        self.profile = import_excel((ROOT/'examples/example-profile.xlsx').read_bytes(), 'Sheet1')

    def test_original_import(self):
        p = self.profile
        self.assertEqual(p['name'], 'LenStandardV17')
        self.assertEqual(p['overview'], {'dia_hours':8.5,'glucose_unit':'mmol/L','timezone':'Europe/Amsterdam','extra':{}})
        self.assertEqual([len(p['schedules'][k]) for k in ('ic','isf','basal','target')], [6,5,8,3])
        self.assertEqual(p['schedules']['target'][1], {'time':'02:00','low':6.1,'high':6.1})
        self.assertEqual(p['schedules']['ic'][3], {'time':'15:00','value':9.5})
        self.assertAlmostEqual(daily_basal(p['schedules']['basal']),18.0)

    def test_link_persists_on_save_and_overwrite_and_rejects_invalid_target(self):
        with tempfile.TemporaryDirectory() as folder:
            standard = save_profile(folder, self.profile)
            draft = deepcopy(standard)
            draft['category'] = 'Drinking'
            draft['linked_profile_id'] = standard['id']
            saved = save_profile(folder, draft)
            loaded = next(p for p in load_profiles(folder)[0] if p['id'] == saved['id'])
            self.assertEqual(loaded['linked_profile_id'], standard['id'])
            draft['linked_profile_id'] = saved['id']
            with self.assertRaisesRegex(ValueError, 'own corresponding'):
                overwrite_profile(folder, draft, loaded)
            draft['linked_profile_id'] = 'missing'
            with self.assertRaisesRegex(ValueError, 'unavailable'):
                save_profile(folder, draft)
            draft['linked_profile_id'] = None
            changed = overwrite_profile(folder, draft, loaded)
            self.assertIsNone(changed['linked_profile_id'])

    def test_change_summary_includes_metadata_timing_and_aligned_values(self):
        before = deepcopy(self.profile)
        after = deepcopy(before)
        after['notes'] = 'Reason for change'
        after['overview']['extra']['Activity'] = 'Test'
        after['schedules']['basal'] = adjust_range(before['schedules']['basal'], 'basal', '05:30', '09:15', 10)
        metadata, schedules = profile_changes(before, after)
        self.assertEqual({r['Setting'] for r in metadata}, {'Notes', 'Overview · Activity', 'Basal start times'})
        self.assertTrue(all(r['Metric'] == 'Basal' for r in schedules))
        self.assertEqual(schedules[0]['From'], '05:30')
        self.assertEqual(schedules[-1]['To'], '09:15')
        self.assertAlmostEqual(sum(r['Change'] * (minute(r['To'],allow_end=True)-minute(r['From']))/60 for r in schedules), .25)
        self.assertEqual(profile_changes(before, before), ([], []))

    def test_excel_export_round_trip_short_long_and_literal_text(self):
        from io import BytesIO
        from openpyxl import load_workbook
        for long_schedule in (False, True):
            with self.subTest(long_schedule=long_schedule):
                p = deepcopy(self.profile)
                p['category'] = 'Drinking / Exercise'
                p['name'] = '=Literal profile name'
                p['notes'] = '=Literal note'
                p['effective_date'] = '2026-09-17'
                p['version'] = 42
                p['overview']['glucose_unit'] = 'mg/dL'
                p['schedules']['basal'] = ([{'time':f'{h:02}:00','value':h/100} for h in range(24)]
                                           if long_schedule else [{'time':'00:00','value':0}])
                content = export_excel(p, ROOT/'examples/profile-import.xlsx')
                sheet = excel_sheets(content)[0]
                imported = import_excel(content, sheet)
                for field in ('name', 'category', 'notes', 'effective_date', 'overview', 'schedules'):
                    self.assertEqual(imported[field], p[field])
                book = load_workbook(BytesIO(content))
                self.assertEqual(book.active['C5'].data_type, 's')
                self.assertEqual(book.active['I5'].data_type, 's')
                self.assertEqual(book.active['L13'].value, '6.6 - 6.6')
                self.assertNotIn('/', sheet)
                self.assertTrue(book.active['F12'].value.startswith('='))
                book.close()

    def test_overwrite_preserves_identity_and_rejects_stale_edits(self):
        with tempfile.TemporaryDirectory() as directory:
            original = save_profile(directory, self.profile)
            path = Path(directory) / f"{original['id']}.json"
            path = path.rename(Path(directory) / 'restored-profile.json')
            draft = deepcopy(original)
            draft['name'] = 'Corrected name'
            draft['overview']['dia_hours'] = 9
            saved = overwrite_profile(directory, draft, original)
            self.assertEqual(len(list(Path(directory).glob('*.json'))), 1)
            self.assertEqual(json.loads(path.read_text()), saved)
            for key in ('id', 'version', 'created_at', 'category', 'schedules'):
                self.assertEqual(saved[key], original[key])
            self.assertEqual(saved['name'], 'Corrected name')
            self.assertIn('updated_at', saved)
            with self.assertRaisesRegex(ValueError, 'changed since'):
                overwrite_profile(directory, draft, original)
            baseline = path.read_bytes()
            with patch('profiles.os.replace', side_effect=OSError('disk failure')):
                with self.assertRaises(OSError):
                    overwrite_profile(directory, draft, saved)
            self.assertEqual(path.read_bytes(), baseline)
            self.assertFalse(list(Path(directory).glob('*.tmp')))
            draft['category'] = 'Other'
            with self.assertRaisesRegex(ValueError, 'Keep the category'):
                overwrite_profile(directory, draft, saved)

    def test_labelled_template_preserves_every_schedule_and_overview_value(self):
        p = import_excel((ROOT/'examples/profile-import.xlsx').read_bytes(), 'Standard-v17')
        for field in ('name', 'category', 'overview', 'schedules', 'effective_date', 'notes'):
            self.assertEqual(p[field], self.profile[field])
        self.assertEqual(p['source']['layout'], 'labelled-v1')

    def labelled_rows(self):
        # Read the actual exported workbook; in-memory edits below exercise the
        # parser without writing alternate workbook fixtures.
        from openpyxl import load_workbook
        book = load_workbook(ROOT/'examples/profile-import.xlsx',read_only=True,data_only=True)
        try:
            return [list(row) for row in book['Standard-v17'].iter_rows(values_only=True)]
        finally:
            book.close()

    def parse_rows(self, rows, sheet='Selected profile'):
        book = MagicMock()
        worksheet = MagicMock()
        worksheet.iter_rows.return_value = rows
        book.__getitem__.side_effect = {sheet:worksheet}.__getitem__
        with patch('openpyxl.load_workbook',return_value=book):
            result = import_excel(b'',sheet)
        book.__getitem__.assert_called_once_with(sheet)
        book.close.assert_called_once()
        return result

    def test_labelled_metadata_and_selected_sheet_are_respected(self):
        rows = self.labelled_rows()
        rows[4][2] = 'My exercise profile'
        rows[5][2] = 'Exercise'
        rows[5][5] = datetime(2026,9,16)
        rows[4][8] = 'Example note'
        p = self.parse_rows(rows)
        self.assertEqual(p['name'],'My exercise profile')
        self.assertEqual(p['category'],'Exercise')
        self.assertEqual(p['effective_date'],'2026-09-16')
        self.assertEqual(p['notes'],'Example note')
        self.assertEqual(p['overview']['extra'],{})
        self.assertEqual(p['schedules'],self.profile['schedules'])

    def test_labelled_template_rejects_missing_metadata_and_invalid_targets(self):
        for row,col,value in [(5,2,None), (5,5,'not a date'), (5,5,123), (12,11,'6.6 - bad'), (4,1,'Wrong label')]:
            rows = self.labelled_rows()
            rows[row][col] = value
            with self.subTest(row=row,col=col,value=value), self.assertRaises(ValueError):
                self.parse_rows(rows)

    def test_range_edit_preserves_every_minute_outside_range(self):
        before = self.profile['schedules']['basal']
        snapshot = deepcopy(before)
        after = adjust_range(before,'basal','05:30','09:15',10)
        for t in range(1440):
            expected = value_at(before,t) * (1.1 if 330 <= t < 555 else 1)
            self.assertAlmostEqual(value_at(after,t),expected)
        self.assertEqual(before,snapshot)
        self.assertAlmostEqual(daily_basal(after)-daily_basal(before),.25)

    def test_full_day_and_target_ranges(self):
        before = self.profile['schedules']['basal']
        after = adjust_range(before,'basal','00:00','24:00',-100)
        self.assertEqual(daily_basal(after),0)
        target = adjust_range(self.profile['schedules']['target'],'target','01:00','02:00',5)
        self.assertAlmostEqual(value_at(target,60,'low'),6.93)
        self.assertEqual(value_at(target,120,'low'),6.1)
        for percent in (-100,-110,float('nan')):
            with self.assertRaises(ValueError):
                adjust_range(self.profile['schedules']['ic'],'ic','00:00','24:00',percent)
        with self.assertRaises(ValueError):
            adjust_range(before,'basal','23:00','01:00',10)

    def test_comparison_aligns_time_not_rows(self):
        reference = [{'time':'00:00','value':0}, {'time':'06:00','value':1}]
        draft = [{'time':'00:00','value':1}, {'time':'03:00','value':2}]
        result = differences(draft,reference,'basal')
        self.assertEqual([row['From'] for row in result],['00:00','03:00','06:00'])
        self.assertEqual(result[1]['Reference'],0)
        self.assertIsNone(result[1]['Change (%)'])
        self.assertEqual(result[2]['Change (%)'],100)

    def test_validation_rejects_incomplete_or_ambiguous_schedules(self):
        invalid = [[], [{'time':'01:00','value':1}], [{'time':'00:00','value':1},{'time':'00:00','value':2}],
                   [{'time':'00:00','value':float('nan')}], [{'time':'00:00','value':0}],
                   [{'time':'00:00','value':-1}], [{'time':'24:00','value':1}]]
        for rows in invalid:
            with self.assertRaises(ValueError):
                normalize_schedule(rows,'ic')
        with self.assertRaises(ValueError):
            normalize_schedule([{'time':'00:00','low':7,'high':6}],'target')
        self.assertEqual(minute(1/12),120)
        with self.assertRaises(ValueError):
            minute('12:60')

    def test_history_is_immutable_and_category_versions_are_independent(self):
        with tempfile.TemporaryDirectory() as directory:
            first = save_profile(directory,self.profile)
            path = Path(directory)/f"{first['id']}.json"
            original = path.read_bytes()
            draft = deepcopy(first)
            draft['parent_id'] = first['id']
            draft['schedules']['basal'][0]['value'] = .9
            second = save_profile(directory,draft,first['id'])
            self.assertEqual(second['version'],2)
            self.assertEqual(second['comparison_id'],first['id'])
            self.assertEqual(second['parent_id'],first['id'])
            self.assertEqual(path.read_bytes(),original)
            draft['category'] = 'Exercise'
            third = save_profile(directory,draft)
            self.assertEqual(third['version'],1)
            profiles,errors = load_profiles(directory)
            self.assertEqual(len(profiles),3)
            self.assertFalse(errors)
            self.assertEqual(validate_profile(json.loads(json.dumps(second))),second)
            (Path(directory)/'broken.json').write_text('{')
            with self.assertRaises(ValueError):
                save_profile(directory,draft)


if __name__ == '__main__':
    unittest.main()
