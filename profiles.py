"""Profile data, interval calculations, Excel import/export and saved versions."""
from __future__ import annotations

from copy import deepcopy
from datetime import date, datetime, time, timezone
from io import BytesIO
import json
import math
from pathlib import Path
import re
import os
import tempfile
from uuid import uuid4
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

METRICS = {"ic": "I:C", "isf": "ISF", "basal": "Basal", "target": "Target"}


def minute(value, *, allow_end=False):
    """Read a clock time, including Excel day fractions; never interpret as hours."""
    if isinstance(value, (datetime, time)):
        if value.second or value.microsecond:
            raise ValueError("Times must use whole minutes.")
        result = value.hour * 60 + value.minute
    elif isinstance(value, (float, int)) and not isinstance(value, bool):
        if not math.isfinite(value) or not 0 <= value <= 1:
            raise ValueError(f"Invalid Excel time: {value!r}")
        result = round(value * 1440)
        if abs(value * 1440 - result) > 1e-6:
            raise ValueError("Times must use whole minutes.")
    else:
        match = re.fullmatch(r"(\d{1,2}):(\d{2})(?::00)?", str(value).strip())
        if not match or int(match[2]) > 59:
            raise ValueError(f"Invalid time {value!r}; use HH:MM.")
        result = int(match[1]) * 60 + int(match[2])
    if not 0 <= result < 1440 and not (allow_end and result == 1440):
        raise ValueError("Use 00:00–23:59 (24:00 is allowed only as a range end).")
    return result


def clock(value):
    return f"{value // 60:02d}:{value % 60:02d}"


def number(value, label, *, zero=False):
    if isinstance(value, bool):
        raise ValueError(f"{label} must be a number.")
    try:
        result = float(value)
    except (TypeError, ValueError):
        raise ValueError(f"{label} must be a number.") from None
    if not math.isfinite(result) or result < 0 or (not zero and result == 0):
        raise ValueError(f"{label} must be finite and {'non-negative' if zero else 'positive'}.")
    return result


def normalize_schedule(rows, metric):
    if metric not in METRICS or not isinstance(rows, list) or not rows:
        raise ValueError(f"{metric}: add at least the 00:00 entry.")
    output = []
    fields = ("low", "high") if metric == "target" else ("value",)
    for row in rows:
        if not isinstance(row, dict):
            raise ValueError(f"{METRICS[metric]}: each entry must be an object.")
        entry = {"time": clock(minute(row.get("time")))}
        for field in fields:
            entry[field] = number(row.get(field), f"{METRICS[metric]} {entry['time']} {field}", zero=metric == "basal")
        if metric == "target" and entry["low"] > entry["high"]:
            raise ValueError(f"Target {entry['time']}: low must not exceed high.")
        output.append(entry)
    output.sort(key=lambda row: row["time"])
    if output[0]["time"] != "00:00":
        raise ValueError(f"{METRICS[metric]}: the first entry must start at 00:00.")
    if len({row["time"] for row in output}) != len(output):
        raise ValueError(f"{METRICS[metric]}: times must be unique.")
    return output


def validate_profile(profile):
    if not isinstance(profile, dict) or profile.get("schema_version") != 1:
        raise ValueError("Expected a Profile Studio JSON file with schema_version 1.")
    p = deepcopy(profile)
    linked_id = p.setdefault("linked_profile_id", None)
    if linked_id is not None and (not isinstance(linked_id, str) or not linked_id.strip()):
        raise ValueError("Corresponding profile must be a saved profile ID or empty.")
    for field in ("name", "category"):
        if not isinstance(p.get(field), str) or not p[field].strip():
            raise ValueError(f"{field.capitalize()} is required.")
        p[field] = p[field].strip()
    overview = p.get("overview")
    if not isinstance(overview, dict):
        raise ValueError("Overview is missing.")
    overview["dia_hours"] = number(overview.get("dia_hours"), "DIA")
    if overview.get("glucose_unit") not in ("mmol/L", "mg/dL"):
        raise ValueError("Glucose unit must be mmol/L or mg/dL.")
    try:
        ZoneInfo(overview.get("timezone", ""))
    except (ZoneInfoNotFoundError, ValueError, TypeError):
        raise ValueError("Enter a valid timezone, such as Europe/Amsterdam.") from None
    extras = overview.setdefault("extra", {})
    if not isinstance(extras, dict) or any(not isinstance(k, str) or not k.strip() for k in extras):
        raise ValueError("Additional overview fields need nonempty names.")
    if any(not isinstance(v, (str, int, float, bool)) or isinstance(v, float) and not math.isfinite(v) for v in extras.values()):
        raise ValueError("Additional overview values must be text, finite numbers or booleans.")
    if not isinstance(p.get("schedules"), dict):
        raise ValueError("Schedules are missing.")
    for metric in METRICS:
        p["schedules"][metric] = normalize_schedule(p["schedules"].get(metric), metric)
    if p.get("effective_date"):
        try:
            from datetime import date
            date.fromisoformat(p["effective_date"])
        except (ValueError, TypeError):
            raise ValueError("Effective date must be YYYY-MM-DD or empty.") from None
    return p


def value_at(rows, at, field="value"):
    return next(row[field] for row in reversed(rows) if minute(row["time"]) <= at)


def daily_basal(rows):
    bounds = [minute(row["time"]) for row in rows] + [1440]
    return sum(row["value"] * (bounds[i + 1] - bounds[i]) / 60 for i, row in enumerate(rows))


def differences(draft, reference, metric):
    """Compare over the union of breakpoints, not by row index."""
    bounds = sorted({minute(r["time"]) for r in draft + reference} | {1440})
    fields = ("low", "high") if metric == "target" else ("value",)
    result = []
    for start, end in zip(bounds, bounds[1:]):
        for field in fields:
            current, previous = value_at(draft, start, field), value_at(reference, start, field)
            result.append({"From": clock(start), "To": clock(end), "Field": field,
                           "Reference": previous, "Draft": current, "Change": current - previous,
                           "Change (%)": None if previous == 0 else 100 * (current / previous - 1)})
    return result


def adjust_range(rows, metric, start, end, percent):
    """Scale [start, end); insert boundaries so the rest of the day is unchanged."""
    start, end = minute(start), minute(end, allow_end=True)
    if start >= end:
        raise ValueError("End must be after start. Split an overnight change into two ranges.")
    factor = 1 + float(percent) / 100
    if not math.isfinite(factor) or factor < 0 or (metric != "basal" and factor == 0):
        raise ValueError("This percentage would create invalid values.")
    rows = normalize_schedule(rows, metric)
    fields = ("low", "high") if metric == "target" else ("value",)
    bounds = sorted({minute(r["time"]) for r in rows} | {start} | ({end} if end < 1440 else set()))
    output = []
    for at in bounds:
        row = {"time": clock(at)}
        for field in fields:
            old = value_at(rows, at, field)
            row[field] = round(old * factor, 6) if start <= at < end else old
        output.append(row)
    return normalize_schedule(output, metric)


def load_profiles(directory):
    profiles, errors = [], []
    for path in sorted(Path(directory).glob("*.json")):
        try:
            p = validate_profile(json.loads(path.read_text(encoding="utf-8")))
            if not isinstance(p.get("version"), int) or p["version"] < 1 or not p.get("id") or not p.get("created_at"):
                raise ValueError("Saved profile is missing its ID, version or creation time.")
            if any(old["id"] == p["id"] for old in profiles):
                raise ValueError("Duplicate profile ID.")
            profiles.append(p)
        except (ValueError, OSError, TypeError) as exc:
            errors.append(f"{path.name}: {exc}")
    return sorted(profiles, key=lambda p: (p["category"].casefold(), p["version"], p["created_at"])), errors


def next_version(profiles, category):
    return max((p["version"] for p in profiles if p["category"].casefold() == category.strip().casefold()), default=0) + 1


def validate_link(profile, profiles, own_id=None):
    linked_id = profile.get("linked_profile_id")
    if linked_id is not None:
        if linked_id == own_id:
            raise ValueError("A profile cannot be its own corresponding profile.")
        if not any(item["id"] == linked_id for item in profiles):
            raise ValueError("Corresponding profile is unavailable. Choose another profile or None.")


def profile_changes(before, after):
    """Return metadata changes and interval-aligned schedule changes."""
    metadata, schedules = [], []
    def changed(label, old, new):
        if old != new:
            metadata.append({"Setting": label, "Before": old, "After": new})
    for key, label in (("name", "Name"), ("category", "Category"), ("notes", "Notes"),
                       ("effective_date", "Effective date"), ("linked_profile_id", "Corresponding profile")):
        changed(label, before.get(key), after.get(key))
    for key, label in (("dia_hours", "DIA (hours)"), ("timezone", "Timezone"), ("glucose_unit", "Glucose unit")):
        changed(label, before["overview"][key], after["overview"][key])
    old_extra, new_extra = before["overview"].get("extra", {}), after["overview"].get("extra", {})
    for key in sorted(old_extra.keys() | new_extra.keys()):
        changed("Overview · " + key, old_extra.get(key), new_extra.get(key))
    for metric, label in METRICS.items():
        old, new = before["schedules"][metric], after["schedules"][metric]
        changed(label + " start times", ", ".join(r["time"] for r in old), ", ".join(r["time"] for r in new))
        for row in differences(new, old, metric):
            if abs(row["Change"]) > 1e-10:
                unit = {"ic": "g/U", "basal": "U/h", "isf": after["overview"]["glucose_unit"] + "/U",
                        "target": after["overview"]["glucose_unit"]}[metric]
                schedules.append({"Metric": label, "Unit": unit, **row})
    return metadata, schedules


def save_profile(directory, draft, reference_id=None, second_reference_id=None, *, version=None):
    p = validate_profile(draft)
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    profiles, errors = load_profiles(directory)
    if errors:
        raise ValueError("Resolve unreadable history files before saving: " + "; ".join(errors))
    validate_link(p, profiles)
    p["category"] = next((old["category"] for old in profiles
                          if old["category"].casefold() == p["category"].casefold()), p["category"])
    if version is None:
        version = next_version(profiles, p["category"])
    if isinstance(version, bool) or not isinstance(version, int) or version < 1:
        raise ValueError("Version must be a positive whole number.")
    if any(old["category"].casefold() == p["category"].casefold() and old["version"] == version for old in profiles):
        raise ValueError("That category/version already exists. Use explicit replacement to overwrite it.")
    p.update(id=str(uuid4()), version=version,
             created_at=datetime.now(timezone.utc).isoformat(), comparison_id=reference_id,
             second_comparison_id=second_reference_id)
    p.pop("updated_at", None)
    _write_profile(directory / f"{p['id']}.json", p)
    return p


def overwrite_profile(directory, draft, original, reference_id=None, second_reference_id=None):
    """Replace the loaded version, refusing stale edits or identity changes."""
    p = validate_profile(draft)
    directory = Path(directory)
    profiles, errors = load_profiles(directory)
    if errors:
        raise ValueError("Resolve unreadable history files before saving: " + "; ".join(errors))
    current = next((item for item in profiles if item["id"] == original["id"]), None)
    if current is None:
        raise ValueError("This profile no longer exists. Load an existing profile again.")
    if current != original:
        raise ValueError("This profile changed since you opened it. Download your draft, then reload the profile before overwriting.")
    if p["category"] != current["category"]:
        raise ValueError("Keep the category when overwriting. Copy to a new draft to change category.")
    validate_link(p, profiles, current["id"])
    # Resolve by ID rather than trusting the ID as a filesystem path. Restored
    # history files can have different filenames from their original UUID.
    path = next(path for path in directory.glob("*.json")
                if json.loads(path.read_text(encoding="utf-8")).get("id") == current["id"])
    for field in ("id", "version", "created_at", "parent_id"):
        if field in current:
            p[field] = current[field]
        else:
            p.pop(field, None)
    p.update(updated_at=datetime.now(timezone.utc).isoformat(),
             comparison_id=reference_id, second_comparison_id=second_reference_id)
    _write_profile(path, p)
    return p


def _write_profile(path, p):
    text = json.dumps(p, indent=2, ensure_ascii=False, allow_nan=False) + "\n"
    # Complete the replacement before publishing it, including on Windows.
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=path.parent, suffix=".tmp", delete=False) as handle:
            temporary = Path(handle.name)
            handle.write(text)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        if temporary and temporary.exists():
            temporary.unlink()


def excel_sheets(content):
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        return book.sheetnames
    finally:
        book.close()


def export_excel(profile, template):
    """Export the fixed overview and schedules in the existing import layout.

    Uses the app's existing Python Excel dependency at runtime. JSON remains the
    full-fidelity format for IDs, links and arbitrary additional overview fields.
    """
    from openpyxl import load_workbook
    from decimal import Decimal
    from openpyxl.workbook.properties import CalcProperties
    p = validate_profile(profile)
    book = load_workbook(template)
    try:
        ws = book.active
        title = f"{p['category']}-v{p['version']}" if "version" in p else p["category"]
        ws.title = re.sub(r"[\\/*?:\[\]]", "-", title).strip("'")[:31] or "Profile"
        def put(address, value):
            cell = ws[address]
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"  # Profile names and notes are literal text.
        for address, value in {"C5": p["name"], "C6": p["category"], "C7": p["overview"]["dia_hours"],
                               "C8": p["overview"]["glucose_unit"], "F5": p["overview"]["timezone"],
                               "F6": date.fromisoformat(p["effective_date"]) if p.get("effective_date") else None,
                               "I5": p.get("notes", ""), "H8": "Source: Profile Studio export"}.items():
            put(address, value)
        ws["F6"].number_format = "yyyy-mm-dd"
        for metric, time_col, value_col in (("ic", 2, 3), ("isf", 5, 6), ("basal", 8, 9), ("target", 11, 12)):
            for row in range(13, ws.max_row + 1):
                ws.cell(row, time_col).value = None
                ws.cell(row, value_col).value = None
            time_style = deepcopy(ws.cell(13, time_col)._style)
            value_style = deepcopy(ws.cell(13, value_col)._style)
            for row, entry in enumerate(p["schedules"][metric], 13):
                at = minute(entry["time"])
                cell = ws.cell(row, time_col, time(at // 60, at % 60))
                cell._style = deepcopy(time_style)
                cell.number_format = "hh:mm"
                value = (f"{format(Decimal(str(entry['low'])), 'f')} - {format(Decimal(str(entry['high'])), 'f')}"
                         if metric == "target" else entry["value"])
                cell = ws.cell(row, value_col, value)
                cell._style = deepcopy(value_style)
                if metric != "target":
                    cell.number_format = "0.######"
                ws.row_dimensions[row].height = ws.row_dimensions[13].height
        book.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
        output = BytesIO()
        book.save(output)
        return output.getvalue()
    finally:
        book.close()


def import_excel(content, sheet, category=None):
    """Import one explicitly selected worksheet, using its labelled overview."""
    from openpyxl import load_workbook
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        ws = book[sheet]
        cells = list(ws.iter_rows(values_only=True))
        def cell(row, col):
            return cells[row-1][col-1] if row <= len(cells) and col <= len(cells[row-1]) else None
        if str(cell(2, 2)).strip() == "Diabetes profile":
            fields = {
                "Profile name": (5, 2, 3), "Category": (6, 2, 3),
                "DIA (hours)": (7, 2, 3), "Glucose unit": (8, 2, 3),
                "Timezone": (5, 5, 6), "Effective date": (6, 5, 6), "Notes": (5, 8, 9),
            }
            overview = {}
            for label, (row, label_col, value_col) in fields.items():
                if str(cell(row, label_col)).strip().casefold() != label.casefold():
                    raise ValueError(f"Expected overview label '{label}' at row {row}. Use the included Excel template.")
                overview[label] = cell(row, value_col)
            name = overview["Profile name"]
            profile_category = overview["Category"] if category is None else category
            dia, glucose_unit, zone = overview["DIA (hours)"], overview["Glucose unit"], overview["Timezone"]
            effective = overview["Effective date"]
            if isinstance(effective, datetime):
                effective = effective.date().isoformat()
            elif isinstance(effective, date):
                effective = effective.isoformat()
            elif isinstance(effective, str):
                effective = effective.strip() or None
            elif effective is not None:
                raise ValueError("Effective date must be an Excel date, YYYY-MM-DD text, or blank.")
            notes = str(overview["Notes"]) if overview["Notes"] is not None else ""
            layout = "labelled-v1"
            specs = (("ic",2,3,13), ("isf",5,6,13), ("basal",8,9,13), ("target",11,12,13))
            for metric, time_col, _, _ in specs:
                if str(cell(11, time_col)).strip().casefold() != METRICS[metric].casefold():
                    raise ValueError(f"Expected '{METRICS[metric]}' schedule title on row 11.")
                if str(cell(12, time_col)).strip().casefold() != "start time":
                    raise ValueError(f"{METRICS[metric]}: expected 'Start time' on row 12.")
        elif all(str(cell(1, col)).strip().casefold() == label.casefold()
                 for col, label in ((3, "I:C"), (6, "ISF"), (9, "Basal"))):
            # Retain support for the original file already supplied with the app.
            name = str(cell(1,1) or "Imported profile")
            profile_category = "Standard" if category is None else category
            dia, glucose_unit, zone = cell(4,1), cell(3,1), cell(5,1)
            effective, notes, layout = None, "", "original"
            specs = (("ic",3,4,3), ("isf",6,7,3), ("basal",9,10,3), ("target",12,13,2))
        else:
            raise ValueError("Excel layout not recognized. Use examples/profile-import.xlsx.")
        schedules = {}
        for metric, time_col, value_col, first in specs:
            entries = []
            for row in range(first, len(cells)+1):
                at, value = cell(row, time_col), cell(row, value_col)
                if at is None and value is None:
                    continue
                if at is None or value is None:
                    raise ValueError(f"{METRICS[metric]}, Excel row {row}: time or value is missing.")
                entry = {"time": clock(minute(at))}
                if metric == "target":
                    parts = re.fullmatch(r"\s*(\d+(?:\.\d+)?)\s*[-–]\s*(\d+(?:\.\d+)?)\s*", str(value))
                    if parts:
                        entry.update(low=float(parts[1]), high=float(parts[2]))
                    else:
                        entry.update(low=number(value,"Target"), high=number(value,"Target"))
                else:
                    entry["value"] = number(value, METRICS[metric], zero=metric=="basal")
                entries.append(entry)
            schedules[metric] = entries
        units = {"mmol": "mmol/L", "mmol/l": "mmol/L", "mg/dl": "mg/dL"}
        unit = units.get(str(glucose_unit).strip().lower())
        if not unit:
            raise ValueError("Glucose unit must be mmol, mmol/L or mg/dL.")
        return validate_profile({"schema_version":1, "name":name,
            "category":profile_category, "parent_id":None, "effective_date":effective, "notes":notes,
            "overview":{"dia_hours":dia, "glucose_unit":unit, "timezone":zone, "extra":{}},
            "schedules":schedules, "source":{"format":"xlsx", "sheet":sheet, "layout":layout}})
    finally:
        book.close()
